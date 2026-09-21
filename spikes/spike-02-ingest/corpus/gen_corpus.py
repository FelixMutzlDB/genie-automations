#!/usr/bin/env python3
"""Generate a messy-but-legal xlsx/csv corpus + per-file ground truth.

Expanded after the partner review to include the cases the first corpus was NOT
exercising (so "100%" was over-optimistic): totals/summary rows, German
period-thousands locale, semicolon+cp1252 CSV. Locale/date handling is
per-automation CONFIG (the harness maps each file to its config).

NOT covered here (parser follow-up): 1904-epoch workbooks + serial-60 (covered by
a harness unit test instead), two-tables-on-one-sheet, hidden sheets.
"""
from __future__ import annotations
import json, pathlib, datetime as dt
import openpyxl

OUT = pathlib.Path(__file__).resolve().parent
BASE_ROWS = [
    {"remittance_id": "R1", "invoice_id": "INV1", "amount": "100.00", "pay_date": "2026-01-15"},
    {"remittance_id": "R1", "invoice_id": "INV2", "amount": "1234.56", "pay_date": "2026-02-01"},
    {"remittance_id": "R2", "invoice_id": "INV3", "amount": "9999.99", "pay_date": "2026-03-31"},
]


def _serial(iso):
    return (dt.date.fromisoformat(iso) - dt.date(1899, 12, 30)).days


def truth(name, rows):
    (OUT / f"{name}.truth.json").write_text(json.dumps(rows, indent=2))


def clean_xlsx():
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["remittance_id", "invoice_id", "amount", "pay_date"])
    for r in BASE_ROWS:
        ws.append([r["remittance_id"], r["invoice_id"], float(r["amount"]), dt.date.fromisoformat(r["pay_date"])])
    wb.save(OUT / "clean.xlsx"); truth("clean.xlsx", BASE_ROWS)


def header_offset_alias_xlsx():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Data"
    ws.append(["Receivables Export"]); ws.append(["Generated 2026-04-01"]); ws.append([])
    ws.append(["Remittance", "Invoice No", "Betrag", "Payment Date"])
    for r in BASE_ROWS:
        ws.append([r["remittance_id"], r["invoice_id"], float(r["amount"]), _serial(r["pay_date"])])
    wb.save(OUT / "header_offset_alias.xlsx"); truth("header_offset_alias.xlsx", BASE_ROWS)


def locale_de_xlsx():
    """German locale numbers as TEXT (1.234,56); real date cells; blank decoy row."""
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["remittance_id", "invoice_id", "amount", "pay_date"])
    ws.append(["R1", "INV1", "100,00", dt.date(2026, 1, 15)])
    ws.append([None, None, None, None])
    ws.append(["R1", "INV2", "1.234,56", dt.date(2026, 2, 1)])
    ws.append(["R2", "INV3", "9.999,99", dt.date(2026, 3, 31)])
    wb.save(OUT / "locale_de.xlsx"); truth("locale_de.xlsx", BASE_ROWS)


def totals_row_xlsx():
    """Trailing totals row (keys empty, amount present) must be SKIPPED, not counted."""
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["remittance_id", "invoice_id", "amount", "pay_date"])
    for r in BASE_ROWS:
        ws.append([r["remittance_id"], r["invoice_id"], float(r["amount"]), dt.date.fromisoformat(r["pay_date"])])
    ws.append([None, None, 11334.55, None])          # TOTAL row -> keys empty -> skip
    wb.save(OUT / "totals_row.xlsx"); truth("totals_row.xlsx", BASE_ROWS)


def multisheet_decoy_xlsx():
    wb = openpyxl.Workbook()
    d1 = wb.active; d1.title = "Cover"; d1.append(["notes", "value"]); d1.append(["hello", 1])
    d2 = wb.create_sheet("Summary"); d2.append(["kpi", "q1"]); d2.append(["dso", 42])
    ws = wb.create_sheet("Ledger")
    ws.append(["remittance_id", "invoice_id", "amount", "pay_date"])
    for r in BASE_ROWS:
        ws.append([r["remittance_id"], r["invoice_id"], float(r["amount"]), dt.date.fromisoformat(r["pay_date"])])
    wb.save(OUT / "multisheet_decoy.xlsx"); truth("multisheet_decoy.xlsx", BASE_ROWS)


def clean_csv():
    lines = ["remittance_id,invoice_id,amount,pay_date"]
    for r in BASE_ROWS:
        lines.append(f'{r["remittance_id"]},{r["invoice_id"]},{r["amount"]},{r["pay_date"]}')
    (OUT / "clean.csv").write_text("\n".join(lines) + "\n"); truth("clean.csv", BASE_ROWS)


def en_thousands_csv():
    lines = ["Remittance,Invoice No,Total,Payment Date",
             'R1,INV1,"100.00",2026-01-15', 'R1,INV2,"1,234.56",2026-02-01', 'R2,INV3,"9,999.99",2026-03-31']
    (OUT / "en_thousands.csv").write_text("\n".join(lines) + "\n"); truth("en_thousands.csv", BASE_ROWS)


def german_semicolon_csv():
    """cp1252 + semicolon delimiter + comma-decimal (European export)."""
    lines = ["Remittance;Invoice No;Betrag;Payment Date",
             "R1;INV1;100,00;2026-01-15", "R1;INV2;1.234,56;2026-02-01", "R2;INV3;9.999,99;2026-03-31"]
    (OUT / "german_semicolon.csv").write_bytes(("\n".join(lines) + "\n").encode("cp1252"))
    truth("german_semicolon.csv", BASE_ROWS)


def main():
    for pat in ("*.xlsx", "*.csv", "*.truth.json"):
        for f in OUT.glob(pat):
            f.unlink()
    clean_xlsx(); header_offset_alias_xlsx(); locale_de_xlsx(); totals_row_xlsx()
    multisheet_decoy_xlsx(); clean_csv(); en_thousands_csv(); german_semicolon_csv()
    print("corpus written to", OUT)


if __name__ == "__main__":
    main()
